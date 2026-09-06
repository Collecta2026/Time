import os
import io
import csv
from datetime import datetime, date, timedelta
from calendar import monthrange

from flask import (Flask, render_template, request, redirect, url_for, flash,
                   session, send_file, send_from_directory, abort, Response)
from flask_login import (LoginManager, login_user, logout_user,
                         login_required, current_user)
from jinja2 import ChoiceLoader, FileSystemLoader
from sqlalchemy import func

from models import (db, User, CompanySettings, Department, Grade, SpinePoint,
                    Employee, BenefitPlan, Loan, LoanPayment, Attendance,
                    Leave, PayrollRun, Payslip, DelegationBand, AuditLog,
                    LeaveProvision)
import payroll as pay
from i18n import translate, STRINGS

# --------------------------------------------------------------------------- #
#  App / config
# --------------------------------------------------------------------------- #
app = Flask(__name__, static_folder=None)
app.secret_key = os.environ.get("SECRET_KEY", "time-dev-secret-change-me")

db_url = os.environ.get("DATABASE_URL", "sqlite:///time.db")
if db_url.startswith("postgres://"):
    db_url = db_url.replace("postgres://", "postgresql://", 1)
app.config["SQLALCHEMY_DATABASE_URI"] = db_url
app.config["SQLALCHEMY_ENGINE_OPTIONS"] = {"pool_pre_ping": True}
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False

# Load templates from ./templates OR repo root (survives flattened uploads)
app.jinja_loader = ChoiceLoader([FileSystemLoader("templates"), FileSystemLoader(".")])


@app.route("/static/<path:filename>")
def static(filename):
    """Serve static assets from ./static OR the repo root (survives flat uploads)."""
    for base in ("static", "."):
        if os.path.isfile(os.path.join(base, filename)):
            return send_from_directory(base, filename)
    abort(404)

db.init_app(app)
login_manager = LoginManager(app)
login_manager.login_view = "login"


@login_manager.user_loader
def load_user(uid):
    return db.session.get(User, int(uid))


from functools import wraps


def require_role(*roles):
    """Gate a route to given roles (admin always allowed)."""
    def deco(fn):
        @wraps(fn)
        @login_required
        def wrapper(*a, **kw):
            if not current_user.has_role(*roles):
                flash("You don't have permission for that action.", "error")
                return redirect(request.referrer or url_for("dashboard"))
            return fn(*a, **kw)
        return wrapper
    return deco


def audit(action, target="", detail=""):
    who = getattr(current_user, "name", None) or getattr(current_user, "email", "system")
    db.session.add(AuditLog(actor=who, action=action, target=target, detail=detail))


def required_signoffs(net_total):
    """Which sign-offs the scheme of delegation demands for this net total."""
    band = DelegationBand.query.filter(
        DelegationBand.min_amount <= net_total,
        db.or_(DelegationBand.max_amount == None,           # noqa: E711
               DelegationBand.max_amount >= net_total)
    ).order_by(DelegationBand.order).first()
    if not band:
        return {"fm": True, "md": True, "band": None}
    return {"fm": band.requires_fm, "md": band.requires_md, "band": band}


# --------------------------------------------------------------------------- #
#  i18n + template helpers
# --------------------------------------------------------------------------- #
def current_lang():
    return session.get("lang", "en")


@app.context_processor
def inject_globals():
    lang = current_lang()

    def t(key):
        return translate(key, lang)

    def loc(obj, field):
        """Return the Arabic field if lang==ar and it exists, else the base field."""
        if lang == "ar":
            ar = getattr(obj, field + "_ar", None)
            if ar:
                return ar
        return getattr(obj, field, "") or ""

    from markupsafe import Markup

    def ic(name):
        return Markup(_ICONS.get(name, ""))

    return dict(t=t, loc=loc, ic=ic, lang=lang,
                direction=("rtl" if lang == "ar" else "ltr"),
                settings=CompanySettings.get(), current_year=date.today().year,
                current_month=date.today().month)


_I = ('<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" '
      'stroke-linecap="round" stroke-linejoin="round">{}</svg>')
_ICONS = {
    "grid": _I.format('<rect x="3" y="3" width="7" height="7"/><rect x="14" y="3" width="7" height="7"/><rect x="14" y="14" width="7" height="7"/><rect x="3" y="14" width="7" height="7"/>'),
    "users": _I.format('<path d="M17 21v-2a4 4 0 0 0-4-4H5a4 4 0 0 0-4 4v2"/><circle cx="9" cy="7" r="4"/><path d="M23 21v-2a4 4 0 0 0-3-3.87"/><path d="M16 3.13a4 4 0 0 1 0 7.75"/>'),
    "layers": _I.format('<polygon points="12 2 2 7 12 12 22 7 12 2"/><polyline points="2 17 12 22 22 17"/><polyline points="2 12 12 17 22 12"/>'),
    "folder": _I.format('<path d="M22 19a2 2 0 0 1-2 2H4a2 2 0 0 1-2-2V5a2 2 0 0 1 2-2h5l2 3h9a2 2 0 0 1 2 2z"/>'),
    "clock": _I.format('<circle cx="12" cy="12" r="10"/><polyline points="12 6 12 12 16 14"/>'),
    "sun": _I.format('<circle cx="12" cy="12" r="5"/><line x1="12" y1="1" x2="12" y2="3"/><line x1="12" y1="21" x2="12" y2="23"/><line x1="4.2" y1="4.2" x2="5.6" y2="5.6"/><line x1="18.4" y1="18.4" x2="19.8" y2="19.8"/><line x1="1" y1="12" x2="3" y2="12"/><line x1="21" y1="12" x2="23" y2="12"/><line x1="4.2" y1="19.8" x2="5.6" y2="18.4"/><line x1="18.4" y1="5.6" x2="19.8" y2="4.2"/>'),
    "wallet": _I.format('<path d="M21 12V7H5a2 2 0 0 1 0-4h14v4"/><path d="M3 5v14a2 2 0 0 0 2 2h16v-5"/><path d="M18 12a2 2 0 0 0 0 4h4v-4z"/>'),
    "heart": _I.format('<path d="M20.8 4.6a5.5 5.5 0 0 0-7.8 0L12 5.7l-1-1.1a5.5 5.5 0 0 0-7.8 7.8l1 1L12 21l7.8-7.6 1-1a5.5 5.5 0 0 0 0-7.8z"/>'),
    "cash": _I.format('<rect x="2" y="6" width="20" height="12" rx="2"/><circle cx="12" cy="12" r="3"/><path d="M6 12h.01M18 12h.01"/>'),
    "chart": _I.format('<line x1="18" y1="20" x2="18" y2="10"/><line x1="12" y1="20" x2="12" y2="4"/><line x1="6" y1="20" x2="6" y2="14"/>'),
    "cog": _I.format('<circle cx="12" cy="12" r="3"/><path d="M19.4 15a1.65 1.65 0 0 0 .33 1.82l.06.06a2 2 0 1 1-2.83 2.83l-.06-.06a1.65 1.65 0 0 0-1.82-.33 1.65 1.65 0 0 0-1 1.51V21a2 2 0 0 1-4 0v-.09A1.65 1.65 0 0 0 9 19.4a1.65 1.65 0 0 0-1.82.33l-.06.06a2 2 0 1 1-2.83-2.83l.06-.06a1.65 1.65 0 0 0 .33-1.82 1.65 1.65 0 0 0-1.51-1H3a2 2 0 0 1 0-4h.09A1.65 1.65 0 0 0 4.6 9a1.65 1.65 0 0 0-.33-1.82l-.06-.06a2 2 0 1 1 2.83-2.83l.06.06a1.65 1.65 0 0 0 1.82.33H9a1.65 1.65 0 0 0 1-1.51V3a2 2 0 0 1 4 0v.09a1.65 1.65 0 0 0 1 1.51 1.65 1.65 0 0 0 1.82-.33l.06-.06a2 2 0 1 1 2.83 2.83l-.06.06a1.65 1.65 0 0 0-.33 1.82V9a1.65 1.65 0 0 0 1.51 1H21a2 2 0 0 1 0 4h-.09a1.65 1.65 0 0 0-1.51 1z"/>'),
    "menu": _I.format('<line x1="3" y1="12" x2="21" y2="12"/><line x1="3" y1="6" x2="21" y2="6"/><line x1="3" y1="18" x2="21" y2="18"/>'),
    "plus": _I.format('<line x1="12" y1="5" x2="12" y2="19"/><line x1="5" y1="12" x2="19" y2="12"/>'),
    "download": _I.format('<path d="M21 15v4a2 2 0 0 1-2 2H5a2 2 0 0 1-2-2v-4"/><polyline points="7 10 12 15 17 10"/><line x1="12" y1="15" x2="12" y2="3"/>'),
    "print": _I.format('<polyline points="6 9 6 2 18 2 18 9"/><path d="M6 18H4a2 2 0 0 1-2-2v-5a2 2 0 0 1 2-2h16a2 2 0 0 1 2 2v5a2 2 0 0 1-2 2h-2"/><rect x="6" y="14" width="12" height="8"/>'),
    "upload": _I.format('<path d="M21 15v4a2 2 0 0 1-2 2H5a2 2 0 0 1-2-2v-4"/><polyline points="17 8 12 3 7 8"/><line x1="12" y1="3" x2="12" y2="15"/>'),
}


@app.template_filter("money")
def money(v):
    try:
        return f"{float(v or 0):,.2f}"
    except (TypeError, ValueError):
        return v


@app.template_filter("d")
def datefmt(v):
    return v.strftime("%Y-%m-%d") if v else ""


@app.route("/lang/<code>")
def set_lang(code):
    session["lang"] = "ar" if code == "ar" else "en"
    return redirect(request.referrer or url_for("dashboard"))


# --------------------------------------------------------------------------- #
#  Auth
# --------------------------------------------------------------------------- #
@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        pw = request.form.get("password", "")
        user = User.query.filter_by(email=email).first()
        if user and user.active and user.check_password(pw):
            login_user(user)
            return redirect(url_for("dashboard"))
        flash("Invalid credentials", "error")
    return render_template("login.html")


@app.route("/logout")
@login_required
def logout():
    logout_user()
    return redirect(url_for("login"))


@app.route("/help")
@login_required
def help_page():
    return render_template("help.html")


@app.route("/healthz")
def healthz():
    return "ok", 200


# --------------------------------------------------------------------------- #
#  Dashboard
# --------------------------------------------------------------------------- #
@app.route("/")
@login_required
def dashboard():
    active = Employee.query.filter_by(status="active").all()
    headcount = len(active)
    gross = sum(e.total_monthly for e in active)

    # Employer cost = gross + employer SI + emergency fund
    s = CompanySettings.get()
    er_cost = 0
    for e in active:
        iw = pay.insurable_wage(e.insurable_wage or e.total_monthly, s.si_floor, s.si_ceiling)
        si = pay.social_insurance(iw, s)
        er_cost += e.total_monthly + float(si["employer_total"])

    outstanding = db.session.query(func.coalesce(func.sum(Loan.outstanding), 0)) \
        .filter(Loan.status == "active").scalar()

    y = date.today().year
    starters = Employee.query.filter(func.extract("year", Employee.hire_date) == y).count()
    leavers = Employee.query.filter(func.extract("year", Employee.end_date) == y).count()

    dept_rows = db.session.query(Department.name, func.count(Employee.id)) \
        .outerjoin(Employee, (Employee.department_id == Department.id) &
                   (Employee.status == "active")) \
        .group_by(Department.id).all()

    recent = Employee.query.filter_by(status="active") \
        .order_by(Employee.hire_date.desc().nullslast()).limit(6).all()

    return render_template("dashboard.html", headcount=headcount, gross=gross,
                           er_cost=er_cost, outstanding=outstanding,
                           starters=starters, leavers=leavers,
                           dept_rows=dept_rows, recent=recent)


# --------------------------------------------------------------------------- #
#  Employees
# --------------------------------------------------------------------------- #
@app.route("/employees")
@login_required
def employees():
    q = request.args.get("q", "").strip()
    status = request.args.get("status", "active")
    query = Employee.query
    if status in ("active", "left"):
        query = query.filter_by(status=status)
    if q:
        like = f"%{q}%"
        query = query.filter(db.or_(Employee.name.ilike(like),
                                    Employee.name_ar.ilike(like),
                                    Employee.code.ilike(like)))
    rows = query.order_by(Employee.code).all()
    return render_template("employees.html", rows=rows, q=q, status=status)


@app.route("/employees/new", methods=["GET", "POST"])
@app.route("/employees/<int:eid>/edit", methods=["GET", "POST"])
@login_required
def employee_form(eid=None):
    e = db.session.get(Employee, eid) if eid else Employee()
    if request.method == "POST":
        f = request.form
        e.code = f.get("code", "").strip()
        e.name = f.get("name", "").strip()
        e.name_ar = f.get("name_ar", "").strip()
        e.national_id = f.get("national_id", "").strip()
        e.gender = f.get("gender", "")
        e.dob = _parse_date(f.get("dob"))
        e.phone = f.get("phone", "")
        e.email = f.get("email", "")
        e.address = f.get("address", "")
        e.department_id = _int(f.get("department_id"))
        e.job_title = f.get("job_title", "")
        e.job_title_ar = f.get("job_title_ar", "")
        e.grade_id = _int(f.get("grade_id"))
        e.spine_point = _int(f.get("spine_point"))
        e.hire_date = _parse_date(f.get("hire_date"))
        e.end_date = _parse_date(f.get("end_date"))
        e.status = f.get("status", "active")
        e.contract_type = f.get("contract_type", "indefinite")
        e.basic_salary = _float(f.get("basic_salary"))
        e.allowances = _float(f.get("allowances"))
        e.insurable_wage = _float(f.get("insurable_wage")) or None
        e.bank_name = f.get("bank_name", "")
        e.bank_account = f.get("bank_account", "")
        e.benefit_id = _int(f.get("benefit_id"))
        if not e.id:
            db.session.add(e)
        db.session.commit()
        flash(translate("saved", current_lang()), "ok")
        return redirect(url_for("employee_detail", eid=e.id))
    return render_template("employee_form.html", e=e,
                           departments=Department.query.order_by(Department.name).all(),
                           grades=Grade.query.order_by(Grade.order).all(),
                           benefits=BenefitPlan.query.filter_by(active=True).all())


@app.route("/employees/<int:eid>")
@login_required
def employee_detail(eid):
    e = db.session.get(Employee, eid) or abort(404)
    payslips = Payslip.query.filter_by(employee_id=eid) \
        .join(PayrollRun).order_by(PayrollRun.year.desc(), PayrollRun.month.desc()).all()
    leaves = Leave.query.filter_by(employee_id=eid).order_by(Leave.start_date.desc()).all()
    # spine headroom
    headroom = []
    if e.grade:
        headroom = [p for p in e.grade.points if e.spine_point and p.point_no > e.spine_point]
    return render_template("employee_detail.html", e=e, payslips=payslips,
                           leaves=leaves, headroom=headroom)


@app.route("/employees/<int:eid>/delete", methods=["POST"])
@login_required
def employee_delete(eid):
    e = db.session.get(Employee, eid) or abort(404)
    db.session.delete(e)
    db.session.commit()
    return redirect(url_for("employees"))


# --------------------------------------------------------------------------- #
#  Departments
# --------------------------------------------------------------------------- #
@app.route("/departments", methods=["GET", "POST"])
@login_required
def departments():
    if request.method == "POST":
        d = Department(name=request.form.get("name", "").strip(),
                       name_ar=request.form.get("name_ar", "").strip())
        if d.name:
            db.session.add(d)
            db.session.commit()
        return redirect(url_for("departments"))
    rows = Department.query.order_by(Department.name).all()
    counts = {d.id: len([e for e in d.employees if e.status == "active"]) for d in rows}
    return render_template("departments.html", rows=rows, counts=counts)


@app.route("/departments/<int:did>/delete", methods=["POST"])
@login_required
def department_delete(did):
    d = db.session.get(Department, did) or abort(404)
    db.session.delete(d)
    db.session.commit()
    return redirect(url_for("departments"))


# --------------------------------------------------------------------------- #
#  Grades & spine points
# --------------------------------------------------------------------------- #
@app.route("/grades")
@login_required
def grades():
    return render_template("grades.html", rows=Grade.query.order_by(Grade.order).all())


@app.route("/grades/new", methods=["GET", "POST"])
@app.route("/grades/<int:gid>/edit", methods=["GET", "POST"])
@login_required
def grade_form(gid=None):
    g = db.session.get(Grade, gid) if gid else Grade()
    if request.method == "POST":
        g.code = request.form.get("code", "").strip()
        g.name = request.form.get("name", "").strip()
        g.name_ar = request.form.get("name_ar", "").strip()
        g.order = _int(request.form.get("order")) or 0
        if not g.id:
            db.session.add(g)
        db.session.commit()
        return redirect(url_for("grade_form", gid=g.id))
    return render_template("grade_form.html", g=g)


@app.route("/grades/<int:gid>/point", methods=["POST"])
@login_required
def grade_add_point(gid):
    g = db.session.get(Grade, gid) or abort(404)
    p = SpinePoint(grade_id=g.id, point_no=_int(request.form.get("point_no")) or 0,
                   monthly_salary=_float(request.form.get("monthly_salary")))
    db.session.add(p)
    db.session.commit()
    return redirect(url_for("grade_form", gid=gid))


@app.route("/points/<int:pid>/delete", methods=["POST"])
@login_required
def point_delete(pid):
    p = db.session.get(SpinePoint, pid) or abort(404)
    gid = p.grade_id
    db.session.delete(p)
    db.session.commit()
    return redirect(url_for("grade_form", gid=gid))


@app.route("/grades/<int:gid>/delete", methods=["POST"])
@login_required
def grade_delete(gid):
    g = db.session.get(Grade, gid) or abort(404)
    db.session.delete(g)
    db.session.commit()
    return redirect(url_for("grades"))


# --------------------------------------------------------------------------- #
#  Loans
# --------------------------------------------------------------------------- #
@app.route("/loans")
@login_required
def loans():
    rows = Loan.query.order_by(Loan.status, Loan.issue_date.desc()).all()
    return render_template("loans.html", rows=rows)


@app.route("/loans/new", methods=["GET", "POST"])
@login_required
def loan_form():
    if request.method == "POST":
        principal = _float(request.form.get("principal"))
        loan = Loan(employee_id=_int(request.form.get("employee_id")),
                    issue_date=_parse_date(request.form.get("issue_date")) or date.today(),
                    principal=principal,
                    monthly_deduction=_float(request.form.get("monthly_deduction")),
                    outstanding=principal,
                    reason=request.form.get("reason", ""))
        db.session.add(loan)
        db.session.commit()
        return redirect(url_for("loans"))
    return render_template("loan_form.html",
                           employees=Employee.query.filter_by(status="active")
                           .order_by(Employee.name).all())


@app.route("/loans/<int:lid>/pay", methods=["POST"])
@login_required
def loan_pay(lid):
    loan = db.session.get(Loan, lid) or abort(404)
    amt = _float(request.form.get("amount"))
    if amt > 0:
        loan.outstanding = max(0, (loan.outstanding or 0) - amt)
        db.session.add(LoanPayment(loan_id=loan.id, amount=amt,
                                   note=request.form.get("note", "manual")))
        if loan.outstanding <= 0:
            loan.status = "settled"
        db.session.commit()
    return redirect(url_for("loans"))


# --------------------------------------------------------------------------- #
#  Medical benefits
# --------------------------------------------------------------------------- #
@app.route("/benefits", methods=["GET", "POST"])
@login_required
def benefits():
    if request.method == "POST":
        bid = _int(request.form.get("id"))
        b = db.session.get(BenefitPlan, bid) if bid else BenefitPlan()
        b.name = request.form.get("name", "").strip()
        b.name_ar = request.form.get("name_ar", "").strip()
        b.provider = request.form.get("provider", "")
        b.employer_cost = _float(request.form.get("employer_cost"))
        b.employee_cost = _float(request.form.get("employee_cost"))
        b.coverage = request.form.get("coverage", "")
        b.active = True
        if not b.id:
            db.session.add(b)
        db.session.commit()
        return redirect(url_for("benefits"))
    return render_template("benefits.html",
                           rows=BenefitPlan.query.order_by(BenefitPlan.name).all())


@app.route("/benefits/<int:bid>/delete", methods=["POST"])
@login_required
def benefit_delete(bid):
    b = db.session.get(BenefitPlan, bid) or abort(404)
    b.active = False
    db.session.commit()
    return redirect(url_for("benefits"))


# --------------------------------------------------------------------------- #
#  Attendance (imported from external fingerprint system)
# --------------------------------------------------------------------------- #
@app.route("/attendance")
@login_required
def attendance():
    today = date.today()
    y = _int(request.args.get("year")) or today.year
    m = _int(request.args.get("month")) or today.month
    start, end = _month_bounds(y, m)
    rows = Attendance.query.filter(Attendance.work_date >= start,
                                   Attendance.work_date <= end) \
        .order_by(Attendance.work_date.desc()).limit(500).all()
    return render_template("attendance.html", rows=rows, y=y, m=m)


@app.route("/attendance/template")
@login_required
def attendance_template():
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["employee_code", "date", "hours_worked", "time_in", "time_out"])
    w.writerow(["1001", "2026-01-04", "8", "", ""])
    w.writerow(["1002", "2026-01-04", "", "08:30", "17:00"])
    return Response(out.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition":
                             "attachment; filename=attendance_template.csv"})


@app.route("/attendance/import", methods=["GET", "POST"])
@login_required
def attendance_import():
    if request.method == "POST":
        file = request.files.get("file")
        if not file:
            flash("No file", "error")
            return redirect(url_for("attendance_import"))
        s = CompanySettings.get()
        text = file.read().decode("utf-8-sig", errors="ignore")
        reader = csv.DictReader(io.StringIO(text))
        emap = {e.code: e for e in Employee.query.all()}
        added = updated = skipped = 0
        for row in reader:
            code = (row.get("employee_code") or row.get("code") or "").strip()
            emp = emap.get(code)
            d = _parse_date((row.get("date") or "").strip())
            if not emp or not d:
                skipped += 1
                continue
            hours = _float(row.get("hours_worked"))
            if not hours and row.get("time_in") and row.get("time_out"):
                hours = _hours_between(row["time_in"], row["time_out"])
            expected = s.standard_daily_hours
            ot_day = max(0.0, hours - expected)
            rec = Attendance.query.filter_by(employee_id=emp.id, work_date=d).first()
            if not rec:
                rec = Attendance(employee_id=emp.id, work_date=d)
                db.session.add(rec)
                added += 1
            else:
                updated += 1
            rec.hours_worked = hours
            rec.expected_hours = expected
            rec.ot_day_hours = ot_day
            rec.status = "present" if hours > 0 else "absent"
        db.session.commit()
        flash(f"Imported {added} new, {updated} updated, {skipped} skipped.", "ok")
        return redirect(url_for("attendance"))
    return render_template("attendance_import.html")


# --------------------------------------------------------------------------- #
#  Leave
# --------------------------------------------------------------------------- #
@app.route("/leave")
@login_required
def leave():
    rows = Leave.query.order_by(Leave.start_date.desc()).limit(300).all()
    return render_template("leave.html", rows=rows)


@app.route("/leave/new", methods=["GET", "POST"])
@login_required
def leave_form():
    if request.method == "POST":
        start = _parse_date(request.form.get("start_date"))
        end = _parse_date(request.form.get("end_date")) or start
        days = (end - start).days + 1 if start and end else 0
        lv = Leave(employee_id=_int(request.form.get("employee_id")),
                   leave_type=request.form.get("leave_type", "annual"),
                   start_date=start, end_date=end, days=days,
                   paid_pct=_float(request.form.get("paid_pct")) or 100,
                   note=request.form.get("note", ""))
        db.session.add(lv)
        db.session.commit()
        return redirect(url_for("leave"))
    return render_template("leave_form.html",
                           employees=Employee.query.filter_by(status="active")
                           .order_by(Employee.name).all())


@app.route("/leave/<int:lid>/delete", methods=["POST"])
@login_required
def leave_delete(lid):
    lv = db.session.get(Leave, lid) or abort(404)
    db.session.delete(lv)
    db.session.commit()
    return redirect(url_for("leave"))


# --------------------------------------------------------------------------- #
#  Payroll
# --------------------------------------------------------------------------- #
@app.route("/payroll")
@login_required
def payroll_runs():
    runs = PayrollRun.query.order_by(PayrollRun.year.desc(),
                                     PayrollRun.month.desc()).all()
    return render_template("payroll.html", runs=runs)


@app.route("/payroll/new", methods=["POST"])
@require_role("hr")
def payroll_new():
    y = _int(request.form.get("year")) or date.today().year
    m = _int(request.form.get("month")) or date.today().month
    existing = PayrollRun.query.filter_by(year=y, month=m).first()
    if existing:
        flash("A run already exists for that period.", "error")
        return redirect(url_for("payroll_detail", rid=existing.id))
    run = PayrollRun(year=y, month=m, status="draft")
    db.session.add(run)
    db.session.commit()
    _compute_run(run)
    audit("payroll.create", run.period_label, f"{len(run.payslips)} payslips")
    db.session.commit()
    return redirect(url_for("payroll_detail", rid=run.id))


def _compute_run(run):
    """Build a payslip for every employee active during the period."""
    s = CompanySettings.get()
    start, end = _month_bounds(run.year, run.month)
    # wipe existing draft slips
    Payslip.query.filter_by(run_id=run.id).delete()

    emps = Employee.query.filter(
        Employee.hire_date <= end,
        db.or_(Employee.end_date == None, Employee.end_date >= start)  # noqa: E711
    ).all()

    for e in emps:
        att = Attendance.query.filter(Attendance.employee_id == e.id,
                                      Attendance.work_date >= start,
                                      Attendance.work_date <= end).all()
        worked = sum(a.hours_worked or 0 for a in att)
        expected = sum(a.expected_hours or 0 for a in att)
        ot_day = sum(max(0.0, (a.hours_worked or 0) - (a.expected_hours or 0)) for a in att)
        ot_night = sum(a.ot_night_hours or 0 for a in att)
        shortfall = sum(max(0.0, (a.expected_hours or 0) - (a.hours_worked or 0)) for a in att)

        loan_ded = sum(min(l.monthly_deduction or 0, l.outstanding or 0)
                       for l in e.active_loans())
        benefit_ded = e.benefit.employee_cost if e.benefit else 0

        r = pay.compute_payslip(
            basic=e.basic_salary, allowances=e.allowances,
            ot_day_hours=ot_day, ot_night_hours=ot_night,
            hours_short=shortfall, loan_deduction=loan_ded,
            other_deductions=benefit_ded,
            insurable_override=e.insurable_wage, s=s)

        db.session.add(Payslip(
            run_id=run.id, employee_id=e.id,
            basic=float(r["basic"]), allowances=float(r["allowances"]),
            overtime=float(r["overtime"]), gross=float(r["gross"]),
            insurable_wage=float(r["insurable_wage"]),
            si_employee=float(r["si_employee"]), si_employer=float(r["si_employer"]),
            emergency_fund=float(r["emergency_fund"]),
            income_tax=float(r["income_tax"]),
            absence_deduction=float(r["absence_deduction"]),
            loan_deduction=float(r["loan_deduction"]),
            other_deductions=float(r["other_deductions"]),
            total_deductions=float(r["total_deductions"]),
            net_pay=float(r["net_pay"]), employer_cost=float(r["employer_cost"]),
            hours_worked=worked, hours_expected=expected))
    db.session.commit()


@app.route("/payroll/<int:rid>")
@login_required
def payroll_detail(rid):
    run = db.session.get(PayrollRun, rid) or abort(404)
    slips = Payslip.query.filter_by(run_id=rid).join(Employee) \
        .order_by(Employee.code).all()
    totals = {
        "gross": sum(p.gross for p in slips),
        "si_employee": sum(p.si_employee for p in slips),
        "si_employer": sum(p.si_employer + p.emergency_fund for p in slips),
        "tax": sum(p.income_tax for p in slips),
        "deductions": sum(p.total_deductions for p in slips),
        "net": sum(p.net_pay for p in slips),
        "employer_cost": sum(p.employer_cost for p in slips),
    }
    signoffs = required_signoffs(totals["net"])
    return render_template("payroll_run.html", run=run, slips=slips, totals=totals,
                           signoffs=signoffs)


@app.route("/payroll/<int:rid>/recompute", methods=["POST"])
@require_role("hr")
def payroll_recompute(rid):
    run = db.session.get(PayrollRun, rid) or abort(404)
    if run.locked:
        flash("Only a draft run can be recomputed.", "error")
    else:
        _compute_run(run)
        audit("payroll.recompute", run.period_label)
        db.session.commit()
        flash("Recomputed.", "ok")
    return redirect(url_for("payroll_detail", rid=rid))


@app.route("/payroll/<int:rid>/submit", methods=["POST"])
@require_role("hr")
def payroll_submit(rid):
    run = db.session.get(PayrollRun, rid) or abort(404)
    if run.status != "draft":
        flash("This run has already been submitted.", "error")
        return redirect(url_for("payroll_detail", rid=rid))
    run.status = "prepared"
    run.prepared_by = current_user.name or current_user.email
    run.prepared_at = datetime.utcnow()
    run.reject_reason = None
    audit("payroll.submit", run.period_label, "submitted for approval")
    db.session.commit()
    flash("Submitted to the Finance Manager for approval.", "ok")
    return redirect(url_for("payroll_detail", rid=rid))


@app.route("/payroll/<int:rid>/approve", methods=["POST"])
@require_role("finance_manager")
def payroll_approve(rid):
    run = db.session.get(PayrollRun, rid) or abort(404)
    if run.status != "prepared":
        flash("This run is not awaiting approval.", "error")
        return redirect(url_for("payroll_detail", rid=rid))
    actor = current_user.name or current_user.email
    # segregation of duties: the approver cannot be the preparer
    if run.prepared_by and actor == run.prepared_by and current_user.role != "admin":
        flash("Segregation of duties: the preparer cannot approve the same run.", "error")
        return redirect(url_for("payroll_detail", rid=rid))
    run.approved_by = actor
    run.approved_at = datetime.utcnow()
    signoffs = required_signoffs(run.net_total)
    if signoffs["md"]:
        run.status = "approved"          # awaiting MD joint release
        msg = "Approved. Awaiting Managing Director authorisation."
    else:
        run.status = "authorised"        # FM sign-off is sufficient in this band
        _apply_loan_deductions(run)
        msg = "Approved and authorised for payment."
    audit("payroll.approve", run.period_label, f"FM approval by {actor}")
    db.session.commit()
    flash(msg, "ok")
    return redirect(url_for("payroll_detail", rid=rid))


@app.route("/payroll/<int:rid>/authorise", methods=["POST"])
@require_role("md")
def payroll_authorise(rid):
    run = db.session.get(PayrollRun, rid) or abort(404)
    if run.status != "approved":
        flash("This run is not awaiting authorisation.", "error")
        return redirect(url_for("payroll_detail", rid=rid))
    run.authorised_by = current_user.name or current_user.email
    run.authorised_at = datetime.utcnow()
    run.status = "authorised"
    _apply_loan_deductions(run)
    audit("payroll.authorise", run.period_label,
          f"MD joint release by {run.authorised_by}")
    db.session.commit()
    flash("Payroll authorised for payment (joint release complete).", "ok")
    return redirect(url_for("payroll_detail", rid=rid))


@app.route("/payroll/<int:rid>/reject", methods=["POST"])
@require_role("finance_manager", "md")
def payroll_reject(rid):
    run = db.session.get(PayrollRun, rid) or abort(404)
    if run.status not in ("prepared", "approved"):
        flash("Only a submitted or approved run can be returned.", "error")
        return redirect(url_for("payroll_detail", rid=rid))
    run.reject_reason = request.form.get("reason", "").strip() or "Returned for review"
    run.status = "draft"
    run.prepared_by = run.prepared_at = None
    run.approved_by = run.approved_at = None
    audit("payroll.reject", run.period_label, run.reject_reason)
    db.session.commit()
    flash("Returned to HR as draft.", "ok")
    return redirect(url_for("payroll_detail", rid=rid))


@app.route("/payroll/<int:rid>/paid", methods=["POST"])
@require_role("finance_manager")
def payroll_mark_paid(rid):
    run = db.session.get(PayrollRun, rid) or abort(404)
    if run.status != "authorised":
        flash("Only an authorised run can be marked paid.", "error")
        return redirect(url_for("payroll_detail", rid=rid))
    run.status = "paid"
    run.paid_by = current_user.name or current_user.email
    run.paid_at = datetime.utcnow()
    audit("payroll.paid", run.period_label, f"marked paid by {run.paid_by}")
    db.session.commit()
    flash("Marked as paid.", "ok")
    return redirect(url_for("payroll_detail", rid=rid))


def _apply_loan_deductions(run):
    """On authorisation, reduce loan balances by this run's deductions (once)."""
    for p in run.payslips:
        if p.loan_deduction and p.employee:
            remaining = p.loan_deduction
            for loan in p.employee.active_loans():
                if remaining <= 0:
                    break
                take = min(remaining, loan.monthly_deduction or 0, loan.outstanding or 0)
                if take > 0:
                    loan.outstanding = max(0, loan.outstanding - take)
                    db.session.add(LoanPayment(loan_id=loan.id, amount=take,
                                               note=f"Payroll {run.period_label}"))
                    if loan.outstanding <= 0:
                        loan.status = "settled"
                    remaining -= take


@app.route("/payroll/<int:rid>/delete", methods=["POST"])
@require_role("hr")
def payroll_delete(rid):
    run = db.session.get(PayrollRun, rid) or abort(404)
    if run.locked:
        flash("Only a draft run can be deleted.", "error")
        return redirect(url_for("payroll_detail", rid=rid))
    audit("payroll.delete", run.period_label)
    db.session.delete(run)
    db.session.commit()
    return redirect(url_for("payroll_runs"))


@app.route("/payroll/<int:rid>/export")
@login_required
def payroll_export(rid):
    run = db.session.get(PayrollRun, rid) or abort(404)
    slips = Payslip.query.filter_by(run_id=rid).join(Employee).order_by(Employee.code).all()
    headers = ["Code", "Employee", "Basic", "Allowances", "Overtime", "Gross",
               "Insurable wage", "SI employee 11%", "Income tax", "Absence",
               "Loan", "Other", "Total deductions", "Net pay",
               "SI employer 18.75%", "Emergency fund 1%", "Employer cost"]
    rows = [[p.employee.code, p.employee.name, p.basic, p.allowances, p.overtime,
             p.gross, p.insurable_wage, p.si_employee, p.income_tax,
             p.absence_deduction, p.loan_deduction, p.other_deductions,
             p.total_deductions, p.net_pay, p.si_employer, p.emergency_fund,
             p.employer_cost] for p in slips]
    return _xlsx_response(f"payroll_{run.period_label}", headers, rows)


@app.route("/payroll/<int:rid>/schedule")
@require_role("hr", "finance_manager", "md")
def payroll_schedule(rid):
    """Bank payment schedule + statutory remittances, for Finance to process."""
    run = db.session.get(PayrollRun, rid) or abort(404)
    slips = Payslip.query.filter_by(run_id=rid).join(Employee).order_by(Employee.code).all()
    totals = {
        "net": sum(p.net_pay for p in slips),
        "tax": sum(p.income_tax for p in slips),
        "si_employee": sum(p.si_employee for p in slips),
        "si_employer": sum(p.si_employer for p in slips),
        "fund": sum(p.emergency_fund for p in slips),
    }
    totals["si_total"] = totals["si_employee"] + totals["si_employer"] + totals["fund"]
    if request.args.get("export"):
        headers = ["Code", "Employee", "Bank", "Account / IBAN", "Net pay (EGP)"]
        rows = [[p.employee.code, p.employee.name, p.employee.bank_name,
                 p.employee.bank_account, p.net_pay] for p in slips]
        rows.append(["", "", "", "TOTAL", totals["net"]])
        return _xlsx_response(f"payment_schedule_{run.period_label}", headers, rows)
    return render_template("payroll_schedule.html", run=run, slips=slips, totals=totals)


@app.route("/payroll/<int:rid>/insurance")
@require_role("hr", "finance_manager", "md")
def payroll_insurance(rid):
    """Auditable social-insurance remittance schedule (NOSI) for one run."""
    run = db.session.get(PayrollRun, rid) or abort(404)
    slips = Payslip.query.filter_by(run_id=rid).join(Employee).order_by(Employee.code).all()
    totals = {
        "insurable": sum(p.insurable_wage for p in slips),
        "employee": sum(p.si_employee for p in slips),
        "employer": sum(p.si_employer for p in slips),
        "fund": sum(p.emergency_fund for p in slips),
    }
    totals["total"] = totals["employee"] + totals["employer"] + totals["fund"]
    if request.args.get("export"):
        headers = ["Code", "Employee", "National ID", "Insurable wage",
                   "Employee 11%", "Employer 18.75%", "Emergency 1%", "Total"]
        rows = [[p.employee.code, p.employee.name, p.employee.national_id,
                 p.insurable_wage, p.si_employee, p.si_employer, p.emergency_fund,
                 p.si_employee + p.si_employer + p.emergency_fund] for p in slips]
        rows.append(["", "", "TOTAL", totals["insurable"], totals["employee"],
                     totals["employer"], totals["fund"], totals["total"]])
        return _xlsx_response(f"insurance_remittance_{run.period_label}", headers, rows)
    return render_template("payroll_insurance.html", run=run, slips=slips, totals=totals)


@app.route("/payslip/<int:pid>")
@login_required
def payslip(pid):
    p = db.session.get(Payslip, pid) or abort(404)
    return render_template("payslip.html", p=p, run=p.run)


# --------------------------------------------------------------------------- #
#  Contract (Arabic, printable)
# --------------------------------------------------------------------------- #
@app.route("/employees/<int:eid>/contract")
@login_required
def contract(eid):
    e = db.session.get(Employee, eid) or abort(404)
    return render_template("contract.html", e=e, s=CompanySettings.get())


# --------------------------------------------------------------------------- #
#  Reports
# --------------------------------------------------------------------------- #
@app.route("/reports")
@login_required
def reports():
    return render_template("reports.html")


@app.route("/reports/attendance")
@login_required
def report_attendance():
    y, m, start, end = _period_from_args()
    rows = db.session.query(
        Employee.code, Employee.name,
        func.coalesce(func.sum(Attendance.hours_worked), 0),
        func.coalesce(func.sum(Attendance.expected_hours), 0),
        func.count(Attendance.id)) \
        .join(Attendance, Attendance.employee_id == Employee.id) \
        .filter(Attendance.work_date >= start, Attendance.work_date <= end) \
        .group_by(Employee.id).order_by(Employee.code).all()
    data = [(c, n, wk, ex, max(0, ex - wk), days) for c, n, wk, ex, days in rows]
    if request.args.get("export"):
        return _xlsx_response(f"attendance_{y}_{m:02d}",
                              ["Code", "Name", "Worked", "Expected", "Shortfall", "Days"],
                              data)
    return render_template("report_attendance.html", data=data, y=y, m=m)


@app.route("/reports/sick")
@login_required
def report_sick():
    y = _int(request.args.get("year")) or date.today().year
    start, end = date(y, 1, 1), date(y, 12, 31)
    rows = db.session.query(
        Employee.code, Employee.name,
        func.coalesce(func.sum(Leave.days), 0), func.count(Leave.id)) \
        .join(Leave, Leave.employee_id == Employee.id) \
        .filter(Leave.leave_type == "sick",
                Leave.start_date >= start, Leave.start_date <= end) \
        .group_by(Employee.id).order_by(func.sum(Leave.days).desc()).all()
    if request.args.get("export"):
        return _xlsx_response(f"sick_{y}", ["Code", "Name", "Sick days", "Episodes"], rows)
    return render_template("report_sick.html", rows=rows, y=y)


@app.route("/reports/starters-leavers")
@login_required
def report_starters_leavers():
    y = _int(request.args.get("year")) or date.today().year
    start, end = date(y, 1, 1), date(y, 12, 31)
    starters = Employee.query.filter(Employee.hire_date >= start,
                                     Employee.hire_date <= end) \
        .order_by(Employee.hire_date).all()
    leavers = Employee.query.filter(Employee.end_date >= start,
                                    Employee.end_date <= end) \
        .order_by(Employee.end_date).all()
    return render_template("report_starters_leavers.html",
                           starters=starters, leavers=leavers, y=y)


@app.route("/reports/stats")
@login_required
def report_stats():
    active = Employee.query.filter_by(status="active").all()
    n = len(active)
    by_dept = db.session.query(Department.name, func.count(Employee.id)) \
        .outerjoin(Employee, (Employee.department_id == Department.id) &
                   (Employee.status == "active")).group_by(Department.id).all()
    by_grade = db.session.query(Grade.code, func.count(Employee.id)) \
        .outerjoin(Employee, (Employee.grade_id == Grade.id) &
                   (Employee.status == "active")).group_by(Grade.id).all()
    males = sum(1 for e in active if e.gender == "male")
    females = sum(1 for e in active if e.gender == "female")
    avg_salary = (sum(e.total_monthly for e in active) / n) if n else 0
    y = date.today().year
    leavers = Employee.query.filter(func.extract("year", Employee.end_date) == y).count()
    turnover = (leavers / n * 100) if n else 0
    return render_template("report_stats.html", n=n, by_dept=by_dept, by_grade=by_grade,
                           males=males, females=females, avg_salary=avg_salary,
                           turnover=turnover, leavers=leavers)


@app.route("/reports/si")
@login_required
def report_si():
    """Yearly social-insurance remittance register — one row per payroll run."""
    y = _int(request.args.get("year")) or date.today().year
    runs = PayrollRun.query.filter_by(year=y).order_by(PayrollRun.month).all()
    reg = []
    for r in runs:
        slips = r.payslips
        reg.append({
            "run": r,
            "headcount": len(slips),
            "insurable": sum(p.insurable_wage for p in slips),
            "employee": sum(p.si_employee for p in slips),
            "employer": sum(p.si_employer for p in slips),
            "fund": sum(p.emergency_fund for p in slips),
            "total": sum(p.si_employee + p.si_employer + p.emergency_fund for p in slips),
        })
    if request.args.get("export"):
        headers = ["Period", "Status", "Headcount", "Insurable total",
                   "Employee 11%", "Employer 18.75%", "Emergency 1%", "Total remittance"]
        rows = [[x["run"].period_label, x["run"].status, x["headcount"], x["insurable"],
                 x["employee"], x["employer"], x["fund"], x["total"]] for x in reg]
        return _xlsx_response(f"insurance_register_{y}", headers, rows)
    return render_template("report_si.html", reg=reg, y=y)


@app.route("/reports/leave-provision", methods=["GET"])
@login_required
def report_leave_provision():
    y, m, _, _ = _period_from_args()
    rows, totals = _compute_leave_provision(y, m)
    posted = LeaveProvision.query.filter_by(year=y).all()
    # movement history: monthly totals of posted snapshots this year
    history = {}
    for p in posted:
        h = history.setdefault(p.month, {"provision": 0, "accrual": 0})
        h["provision"] += p.provision_value
        h["accrual"] += p.monthly_accrual
    history = sorted(history.items())
    is_posted = LeaveProvision.query.filter_by(year=y, month=m).first() is not None
    if request.args.get("export"):
        headers = ["Code", "Employee", "Entitlement (days)", "Accrued (days)",
                   "Taken (days)", "Balance (days)", "Daily rate", "Provision value"]
        data = [[r["emp"].code, r["emp"].name, r["entitlement"], r["accrued"],
                 r["taken"], r["balance"], r["daily"], r["provision"]] for r in rows]
        data.append(["", "", "", "", "", totals["balance_days"], "TOTAL", totals["provision"]])
        return _xlsx_response(f"leave_provision_{y}_{m:02d}", headers, data)
    return render_template("report_leave_provision.html", rows=rows, totals=totals,
                           y=y, m=m, history=history, is_posted=is_posted)


@app.route("/reports/leave-provision/post", methods=["POST"])
@require_role("finance_manager")
def leave_provision_post():
    y = _int(request.form.get("year")) or date.today().year
    m = _int(request.form.get("month")) or date.today().month
    rows, _ = _compute_leave_provision(y, m)
    LeaveProvision.query.filter_by(year=y, month=m).delete()
    who = current_user.name or current_user.email
    for r in rows:
        db.session.add(LeaveProvision(
            year=y, month=m, employee_id=r["emp"].id,
            entitlement_days=r["entitlement"], accrued_days=r["accrued"],
            taken_days=r["taken"], balance_days=r["balance"],
            daily_rate=r["daily"], provision_value=r["provision"],
            monthly_accrual=r["monthly_accrual"], posted_by=who))
    audit("leave_provision.post", f"{y}-{m:02d}",
          f"{len(rows)} employees posted by {who}")
    db.session.commit()
    flash("Leave provision posted for the month.", "ok")
    return redirect(url_for("report_leave_provision", year=y, month=m))


@app.route("/reports/loans")
@login_required
def report_loans():
    rows = Loan.query.join(Employee).order_by(Loan.status, Employee.code).all()
    if request.args.get("export"):
        headers = ["Code", "Name", "Issued", "Principal", "Monthly", "Outstanding", "Status"]
        data = [[l.employee.code, l.employee.name, l.issue_date, l.principal,
                 l.monthly_deduction, l.outstanding, l.status] for l in rows]
        return _xlsx_response("loan_register", headers, data)
    return render_template("report_loans.html", rows=rows)


# --------------------------------------------------------------------------- #
#  Users (administration) + audit log
# --------------------------------------------------------------------------- #
@app.route("/users", methods=["GET", "POST"])
@require_role()   # admin only
def users():
    if request.method == "POST":
        uid = _int(request.form.get("id"))
        u = db.session.get(User, uid) if uid else User()
        u.email = request.form.get("email", "").strip().lower()
        u.name = request.form.get("name", "").strip()
        u.role = request.form.get("role", "viewer")
        u.active = request.form.get("active") == "on"
        pw = request.form.get("password", "")
        if pw:
            u.set_password(pw)
        if not u.id:
            db.session.add(u)
        audit("user.save", u.email, u.role)
        db.session.commit()
        flash(translate("saved", current_lang()), "ok")
        return redirect(url_for("users"))
    return render_template("users.html", rows=User.query.order_by(User.role).all())


@app.route("/users/<int:uid>/delete", methods=["POST"])
@require_role()   # admin only
def user_delete(uid):
    u = db.session.get(User, uid) or abort(404)
    if u.id == current_user.id:
        flash("You can't delete your own account.", "error")
        return redirect(url_for("users"))
    audit("user.delete", u.email)
    db.session.delete(u)
    db.session.commit()
    return redirect(url_for("users"))


@app.route("/audit")
@require_role("finance_manager", "md")
def audit_log():
    rows = AuditLog.query.order_by(AuditLog.ts.desc()).limit(400).all()
    return render_template("audit.html", rows=rows)


# --------------------------------------------------------------------------- #
#  Settings
# --------------------------------------------------------------------------- #
@app.route("/settings", methods=["GET", "POST"])
@require_role()   # admin only
def settings_page():
    s = CompanySettings.get()
    if request.method == "POST":
        f = request.form
        for field in ["company_name", "company_name_ar", "address", "tax_id",
                      "insurance_no", "currency", "working_days",
                      "pay_bank_name", "pay_bank_account"]:
            setattr(s, field, f.get(field, getattr(s, field)))
        for field in ["standard_daily_hours", "standard_weekly_hours",
                      "si_employee_pct", "si_employer_pct", "emergency_fund_pct",
                      "si_floor", "si_ceiling", "tax_annual_exemption",
                      "overtime_day_pct", "overtime_night_pct"]:
            setattr(s, field, _float(f.get(field)))
        for field in ["days_per_month", "annual_leave_y1", "annual_leave_std",
                      "annual_leave_senior", "casual_leave", "probation_months"]:
            setattr(s, field, _int(f.get(field)))
        # scheme of delegation bands (delivered inline)
        for b in DelegationBand.query.all():
            b.min_amount = _float(f.get(f"band_{b.id}_min"))
            b.max_amount = _float(f.get(f"band_{b.id}_max")) or None
            b.requires_fm = f.get(f"band_{b.id}_fm") == "on"
            b.requires_md = f.get(f"band_{b.id}_md") == "on"
        audit("settings.update", "company parameters")
        db.session.commit()
        flash(translate("saved", current_lang()), "ok")
        return redirect(url_for("settings_page"))
    return render_template("settings.html", s=s,
                           bands=DelegationBand.query.order_by(DelegationBand.order).all())


# --------------------------------------------------------------------------- #
#  Helpers
# --------------------------------------------------------------------------- #
def _int(v):
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def _float(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _parse_date(v):
    if not v:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(v.strip(), fmt).date()
        except ValueError:
            continue
    return None


def _hours_between(t_in, t_out):
    try:
        a = datetime.strptime(t_in.strip(), "%H:%M")
        b = datetime.strptime(t_out.strip(), "%H:%M")
        return round((b - a).seconds / 3600, 2)
    except (ValueError, AttributeError):
        return 0.0


def _month_bounds(y, m):
    return date(y, m, 1), date(y, m, monthrange(y, m)[1])


def _period_from_args():
    today = date.today()
    y = _int(request.args.get("year")) or today.year
    m = _int(request.args.get("month")) or today.month
    start, end = _month_bounds(y, m)
    return y, m, start, end


def _leave_entitlement(emp, s):
    """Statutory annual leave entitlement (days) per Law 14/2025 tiers."""
    yrs = emp.years_service or 0
    age = None
    if emp.dob:
        age = (date.today() - emp.dob).days / 365.25
    if yrs >= 10 or (age is not None and age >= 50):
        return s.annual_leave_senior or 30
    if yrs < 1:
        return s.annual_leave_y1 or 15
    return s.annual_leave_std or 21


def _compute_leave_provision(year, month):
    """Live per-employee leave-pay provision to the end of the given month."""
    s = CompanySettings.get()
    start = date(year, 1, 1)
    period_end = date(year, month, monthrange(year, month)[1])
    rows = []
    for e in Employee.query.filter_by(status="active").order_by(Employee.code).all():
        if e.hire_date and e.hire_date > period_end:
            continue
        entitlement = _leave_entitlement(e, s)
        first_month = e.hire_date.month if (e.hire_date and e.hire_date.year == year) else 1
        months_earned = max(0, month - first_month + 1)
        accrued = round(entitlement * months_earned / 12, 2)
        taken = sum(l.days or 0 for l in Leave.query.filter(
            Leave.employee_id == e.id, Leave.leave_type == "annual",
            Leave.start_date >= start, Leave.start_date <= period_end).all())
        balance = round(accrued - taken, 2)
        daily = (e.total_monthly or 0) / (s.days_per_month or 30)
        provision = round(max(0.0, balance) * daily, 2)
        monthly_accrual = round(entitlement / 12 * daily, 2)
        rows.append({"emp": e, "entitlement": entitlement, "accrued": accrued,
                     "taken": taken, "balance": balance, "daily": round(daily, 2),
                     "provision": provision, "monthly_accrual": monthly_accrual})
    totals = {
        "provision": round(sum(r["provision"] for r in rows), 2),
        "monthly_accrual": round(sum(r["monthly_accrual"] for r in rows), 2),
        "balance_days": round(sum(r["balance"] for r in rows), 2),
    }
    return rows, totals


def _xlsx_response(name, headers, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = name[:31]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="101A2B")
    for r in rows:
        ws.append(list(r))
    for i, _ in enumerate(headers, 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"{name}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# --------------------------------------------------------------------------- #
#  Bootstrap
# --------------------------------------------------------------------------- #
def init_db():
    with app.app_context():
        db.create_all()
        CompanySettings.get()
        if not DelegationBand.query.first():
            # default scheme: every payroll needs FM approval + MD joint release
            db.session.add(DelegationBand(name="All payrolls", min_amount=0,
                                          max_amount=None, requires_fm=True,
                                          requires_md=True, order=1))
            db.session.commit()
        if not User.query.first():
            admin = User(email=os.environ.get("ADMIN_EMAIL", "admin@time.eg").lower(),
                         name="Administrator", role="admin")
            admin.set_password(os.environ.get("ADMIN_PASSWORD", "Time2026"))
            db.session.add(admin)
            db.session.commit()
        if os.environ.get("SEED_DEMO", "true").lower() == "true" \
                and Employee.query.count() == 0:
            try:
                import seed
                seed.run(db)
            except Exception as exc:  # pragma: no cover
                print("seed skipped:", exc)


init_db()

if __name__ == "__main__":
    app.run(debug=True, port=5000)
